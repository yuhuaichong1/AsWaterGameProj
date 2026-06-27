# -*- coding: utf-8 -*-
"""Setup Payout Luban config tables and append Language entries."""
import openpyxl
from openpyxl import Workbook
import os

BASE = r'd:\UnityProj\_Project\AsWaterGameProj\DesignerConfigs\Datas'

# --- Language entries: (id, zh, en) ---
LANG_ENTRIES = [
    (10201, "提现资格审核", "Withdrawal Eligibility"),
    (10202, "达成提现金额 {0}/{1}", "Reached payout goal {0}/{1}"),
    (10203, "观看广告 {0}/{1} 次", "Watch video {0}/{1} times"),
    (10204, "为确保提现请求由真实玩家发起，需完成身份核验。请观看 <color=#EA2111>{0}</color> 次视频广告以激活提现资格，有助于保障账户与打款安全。", "To verify your payout request, please watch <color=#EA2111>{0}</color> video ads to activate eligibility and help protect your account."),
    (10205, "收款渠道确认", "Payout Channel Confirmation"),
    (10206, "观看广告 {0}/{1} 次", "Watch video {0}/{1} times"),
    (10207, "完成关卡 {0}/{1} 关", "Complete {0}/{1} levels"),
    (10208, "需确认您的收款渠道处于可用状态且与账户信息一致。请保持在线并完成 <color=#EA2111>{0}</color> 关，有助于减少打款失败并顺利处理提现。", "Please confirm your payout channel is active. Stay online and complete <color=#EA2111>{0}</color> levels to help process your payout smoothly."),
    (10209, "账户活跃验证", "Account Activity Check"),
    (10210, "完成关卡 {0}/{1} 关", "Complete {0}/{1} levels"),
    (10211, "活跃签到 {0}/{1} 天", "Active check-in {0}/{1} days"),
    (10212, "为确认账户已启用且您为真实玩家，请保持连续活跃。今日还需 <color=#EA2111>{0}</color> 关完成签到，不活跃可能导致审核延后。", "Please stay active. Complete <color=#EA2111>{0}</color> more levels today for check-in to avoid review delays."),
    (10213, "风控初审排队", "Risk Review Queue"),
    (10214, "活跃签到 {0}/{1} 天", "Active check-in {0}/{1} days"),
    (10215, "完成关卡 {0}/{1} 关", "Complete {0}/{1} levels"),
    (10216, "您的申请已进入风控复核队列，预计约 <color=#EA2111>{0}</color> 个工作日。请保持账号活跃并完成今日 <color=#EA2111>{1}</color> 关，避免被判定为异常账号而延后处理。", "Your request is in the risk review queue (~<color=#EA2111>{0}</color> business days). Stay active and complete <color=#EA2111>{1}</color> levels today."),
    (10217, "银行通道核验", "Bank Channel Verification"),
    (10218, "风控初审已提交", "Risk review submitted"),
    (10219, "银行审核 {0}/{1} 天", "Bank review {0}/{1} days"),
    (10220, "银行审核约需 <color=#EA2111>{0}</color> 天，请保持账号活跃。不活跃账号可能被插队，导致审核进度延迟。今日还需 <color=#EA2111>{1}</color> 关完成签到。", "Bank review takes ~<color=#EA2111>{0}</color> days. Stay active. Complete <color=#EA2111>{1}</color> levels today."),
    (10221, "手续费合规确认", "Payout Fee Confirmation"),
    (10222, "银行审核 {0}/{1} 天", "Bank review {0}/{1} days"),
    (10223, "完成关卡 {0}/{1} 关", "Complete {0}/{1} levels"),
    (10224, "不同地区支付渠道处理费用可能不同。请保持在线并完成 <color=#EA2111>{0}</color> 关，以便确认适用费率并准确计算到账金额。", "Regional fees may vary. Stay online and complete <color=#EA2111>{0}</color> levels to confirm applicable rates."),
    (10225, "最终打款排队", "Final Payout Queue"),
    (10226, "手续费合规已确认", "Fee compliance confirmed"),
    (10227, "排队进度 {0}/{1}", "Queue progress {0}/{1}"),
    (10228, "当前提现申请较多，将按提交顺序分批处理。请保持账号活跃并完成今日 <color=#EA2111>{0}</color> 关，有助于维持审核顺位。较大金额申请将优先处理，请耐心等待。", "High volume requests are processed in batches. Stay active and complete <color=#EA2111>{0}</color> levels today. Larger amounts are prioritized."),
    (10229, "打款通道排队中", "Payout Channel In Queue"),
    (10230, "最终打款排队已完成", "Final queue step completed"),
    (10231, "保持活跃 {0}/{1} 关", "Stay active {0}/{1} levels"),
    (10232, "当前支付通道较为繁忙，您的申请正在排队处理中。请保持账号活跃并完成每日 <color=#EA2111>{0}</color> 关，以免排队顺位被延后。请耐心等待，无需额外操作。", "The payout channel is busy and your request is queued. Stay active and complete <color=#EA2111>{0}</color> levels daily to maintain your queue position. Please wait patiently."),
    (10233, "剩余 {0} 可继续", "Available in {0}"),
    (10234, "请完成今日任务并等待倒计时结束", "Complete today's tasks and wait for the countdown"),
    (10235, "继续", "Continue"),
]

# Step config rows: sn, waitHours, taskType, taskTarget, onlineMinutes, dailyLevelTarget, titleLang, prevLang, curLang, explainLang, ifTerminal
# taskType: 0=None 1=Amount 2=Ad 3=Level 4=CheckIn 5=BankReview 6=Online 7=Queue 8=DailyLevel
STEPS = [
    (1, 24, 2, 5, 0, 0, 10201, 10202, 10203, 10204, False),
    (2, 24, 3, 5, 5, 0, 10205, 10206, 10207, 10208, False),
    (3, 24, 4, 7, 0, 5, 10209, 10210, 10211, 10212, False),
    (4, 24, 3, 5, 0, 0, 10213, 10214, 10215, 10216, False),
    (5, 24, 5, 2, 0, 5, 10217, 10218, 10219, 10220, False),
    (6, 24, 3, 8, 8, 0, 10221, 10222, 10223, 10224, False),
    (7, 24, 7, 5, 0, 5, 10225, 10226, 10227, 10228, False),
    (8, 24, 8, 5, 0, 5, 10229, 10230, 10231, 10232, True),
]

TIERS = [
    (1, 1000, 1),
    (2, 1500, 2),
    (3, 3000, 3),
    (4, 5000, 4),
    (5, 8000, 5),
    (6, 10000, 6),
]


def append_language():
    path = os.path.join(BASE, 'Language.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb['Language']
    existing = set()
    for row in ws.iter_rows(min_row=6, values_only=True):
        if row[1] is not None:
            try:
                existing.add(int(row[1]))
            except Exception:
                pass
    row_idx = ws.max_row + 1
    added = 0
    for lid, zh, en in LANG_ENTRIES:
        if lid in existing:
            continue
        ws.cell(row=row_idx, column=1, value=None)
        ws.cell(row=row_idx, column=2, value=lid)
        ws.cell(row=row_idx, column=3, value=zh)
        ws.cell(row=row_idx, column=4, value=en)
        for c in range(5, 16):
            ws.cell(row=row_idx, column=c, value=en)
        row_idx += 1
        added += 1
    wb.save(path)
    print(f'Language: appended {added} entries')


def add_enum():
    path = os.path.join(BASE, '__enums__.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    # check if already added
    for r in rows:
        if r[1] == 'item.EPayoutTaskType':
            print('Enum EPayoutTaskType already exists')
            return
    start = ws.max_row + 1
    enum_items = [
        ('item.EPayoutTaskType', 'None', 0),
        (None, 'Amount', 1),
        (None, 'Ad', 2),
        (None, 'Level', 3),
        (None, 'CheckIn', 4),
        (None, 'BankReview', 5),
        (None, 'Online', 6),
        (None, 'Queue', 7),
        (None, 'DailyLevel', 8),
    ]
    first = True
    for full, name, val in enum_items:
        ws.cell(row=start, column=2, value='item.EPayoutTaskType' if first else None)
        ws.cell(row=start, column=3, value=False)
        ws.cell(row=start, column=4, value=True)
        ws.cell(row=start, column=7, value=name)
        ws.cell(row=start, column=9, value=val)
        first = False
        start += 1
    wb.save(path)
    print('Enum EPayoutTaskType added')


def add_tables():
    path = os.path.join(BASE, '__tables__.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    for name, data_file in [('TBPayoutStep', 'PayoutStep.xlsx'), ('TBPayoutTier', 'PayoutTier.xlsx')]:
        exists = any(r[1] == name for r in ws.iter_rows(min_row=4, values_only=True) if r[1])
        if exists:
            print(f'Table {name} already registered')
            continue
        row = ws.max_row + 1
        ws.cell(row=row, column=2, value=name)
        ws.cell(row=row, column=3, value=f'Conf{name[2:]}')
        ws.cell(row=row, column=4, value=True)
        ws.cell(row=row, column=5, value=data_file)
        ws.cell(row=row, column=6, value='sn')
    wb.save(path)
    print('Tables registered')


def create_payout_step_xlsx():
    path = os.path.join(BASE, 'PayoutStep.xlsx')
    if os.path.exists(path):
        print('PayoutStep.xlsx exists, skip create')
        return
    wb = Workbook()
    ws = wb.active
    ws.title = 'PayoutStep'
    headers_var = ['##var', 'sn', 'waitHours', 'taskType', 'taskTarget', 'onlineMinutes', 'dailyLevelTarget',
                   'titleLangId', 'prevTaskLangId', 'curTaskLangId', 'explainLangId', 'ifTerminal']
    headers_type = ['##type', 'int', 'int', 'item.EPayoutTaskType', 'int', 'int', 'int', 'int', 'int', 'int', 'int', 'bool']
    headers_group = ['##group', 'c', 'c', 'c', 'c', 'c', 'c', 'c', 'c', 'c', 'c', 'c']
    headers_desc = ['##', '步骤sn', '等待小时', '任务类型', '任务目标', '在线分钟', '每日通关数', '标题语言id', '上一步语言id', '当前步语言id', '说明语言id', '是否终态']
    ws.append(headers_var)
    ws.append(['##var'] + [None] * (len(headers_var) - 1))
    ws.append(headers_type)
    ws.append(headers_group)
    ws.append(headers_desc)
    for s in STEPS:
        ws.append([None] + list(s))
    wb.save(path)
    print('Created PayoutStep.xlsx')


def create_payout_tier_xlsx():
    path = os.path.join(BASE, 'PayoutTier.xlsx')
    if os.path.exists(path):
        print('PayoutTier.xlsx exists, skip create')
        return
    wb = Workbook()
    ws = wb.active
    ws.title = 'PayoutTier'
    ws.append(['##var', 'sn', 'targetAmount', 'sortOrder'])
    ws.append(['##var', None, None, None])
    ws.append(['##type', 'int', 'float', 'int'])
    ws.append(['##group', 'c', 'c', 'c'])
    ws.append(['##', '档位sn', '目标金额', '排序'])
    for t in TIERS:
        ws.append([None] + list(t))
    wb.save(path)
    print('Created PayoutTier.xlsx')


def add_uires():
    path = os.path.join(BASE, 'UIRes.xlsx')
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    for r in ws.iter_rows(min_row=6, values_only=True):
        if r[1] == 29:
            print('UIRes sn 29 exists')
            break
    else:
        row = ws.max_row + 1
        while ws.cell(row=row - 1, column=2).value is None and row > 6:
            row -= 1
        if ws.cell(row=row, column=2).value is not None:
            row += 1
        ws.cell(row=row, column=1, value=None)
        ws.cell(row=row, column=2, value=29)
        ws.cell(row=row, column=3, value=False)
        ws.cell(row=row, column=4, value=4)
        ws.cell(row=row, column=5, value='Prefabs/UI/ProgressPanel/UIProgressPanel.prefab')
        wb.save(path)
        print('UIRes sn 29 added')
    # fix WithdrawGoal2 path if wrong
    for r in range(6, ws.max_row + 1):
        if ws.cell(row=r, column=2).value == 28:
            path_val = ws.cell(row=r, column=5).value
            if path_val and 'WithdrawGoal/UIWithdrawGoal2' in str(path_val):
                ws.cell(row=r, column=5, value='Prefabs/UI/WithdrawGoal2/UIWithdrawGoal2.prefab')
                wb.save(path)
                print('Fixed UIWithdrawGoal2 path')
            break


if __name__ == '__main__':
    add_enum()
    add_tables()
    create_payout_step_xlsx()
    create_payout_tier_xlsx()
    append_language()
    add_uires()
    print('Done.')
